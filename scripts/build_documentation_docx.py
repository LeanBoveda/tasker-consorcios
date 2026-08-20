from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
OUTPUT = DOCS_DIR / "Documentacion-Tasker-Consorcios.docx"

SOURCE_FILES = [
    DOCS_DIR / "01-manual-de-usuario.md",
    DOCS_DIR / "02-manual-del-administrador.md",
    DOCS_DIR / "03-documentacion-tecnica.md",
    DOCS_DIR / "04-operacion-mantenimiento-y-datos.md",
    DOCS_DIR / "05-referencia-api-y-modelo-de-datos.md",
]

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "173B4D"
MUTED = "66736D"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "C9D3DC"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    paragraph.add_run("Página ")
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Las columnas deben sumar {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))
            cell.width = Inches(widths[index] / 1440)

            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                item = tc_mar.find(qn(f"w:{side}"))
                if item is None:
                    item = OxmlElement(f"w:{side}")
                    tc_mar.append(item)
                item.set(qn("w:w"), str(value))
                item.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def table_widths(headers):
    count = len(headers)
    normalized = [value.lower() for value in headers]
    if count == 2:
        return [2700, 6660]
    if count == 3:
        if normalized[:2] == ["rol visible", "valor interno"]:
            return [2200, 1800, 5360]
        if normalized[:2] == ["columna", "obligatoria"]:
            return [2500, 1500, 5360]
        return [2500, 1450, 5410]
    if count == 4:
        if normalized and normalized[0] == "método":
            return [900, 2700, 3600, 2160]
        return [1500, 2100, 3300, 2460]
    base = CONTENT_WIDTH_DXA // count
    widths = [base] * count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def create_abstract_numbering(doc, bullet=True):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)

    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    level.append(r_pr)
    abstract.append(level)

    first_num_index = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)
    return abstract_id


def create_num_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


def clean_inline_markdown(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"<((?:https?://|mailto:)[^>]+)>", r"\1", text)
    return text


def add_inline_runs(paragraph, text, size=11, color="222222"):
    text = clean_inline_markdown(text)
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=max(size - 1, 8.5), color=DARK_BLUE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def add_markdown_table(doc, rows):
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        add_inline_runs(paragraph, header, size=9.5, color=NAVY)
        for run in paragraph.runs:
            run.bold = True

    for row_values in rows[1:]:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline_runs(paragraph, value, size=9.5)

    set_table_geometry(table, table_widths(headers))
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(0)


def add_code_block(doc, code):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    set_paragraph_shading(paragraph, LIGHT_GRAY)
    lines = code.splitlines() or [""]
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=9, color="253746")
        if index < len(lines) - 1:
            run.add_break()


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        rows.append(values)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", value) for value in rows[1]):
        return [rows[0], *rows[2:]], index
    return None, start


def add_markdown(doc, text, bullet_abstract_id, decimal_abstract_id):
    lines = text.splitlines()
    index = 0
    paragraph_buffer = []
    list_kind = None
    list_num_id = None

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.25
            add_inline_runs(paragraph, " ".join(value.strip() for value in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if line.startswith("```"):
            flush_paragraph()
            list_kind = None
            language = line[3:].strip()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(doc, "\n".join(code_lines))
            index += 1
            continue

        if line.startswith("|"):
            table_rows, next_index = parse_table(lines, index)
            if table_rows:
                flush_paragraph()
                list_kind = None
                add_markdown_table(doc, table_rows)
                index = next_index
                continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            list_kind = None
            level = len(heading.group(1))
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(paragraph, heading.group(2), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE)
            for run in paragraph.runs:
                run.bold = True
            index += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            if list_kind != "bullet":
                list_num_id = create_num_instance(doc, bullet_abstract_id)
            list_kind = "bullet"
            paragraph = doc.add_paragraph()
            apply_numbering(paragraph, list_num_id)
            add_inline_runs(paragraph, bullet.group(1))
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            if list_kind != "decimal":
                list_num_id = create_num_instance(doc, decimal_abstract_id)
            list_kind = "decimal"
            paragraph = doc.add_paragraph()
            apply_numbering(paragraph, list_num_id)
            add_inline_runs(paragraph, numbered.group(1))
            index += 1
            continue

        if not line:
            flush_paragraph()
            list_kind = None
        else:
            list_kind = None
            paragraph_buffer.append(raw)
        index += 1

    flush_paragraph()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("TASKER CONSORCIOS  |  DOCUMENTACIÓN")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_after = Pt(0)
    add_page_field(footer_p)
    for run in footer_p.runs:
        set_run_font(run, size=8.5, color=MUTED)


def add_cover(doc):
    for _ in range(5):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("MANUAL INTEGRAL")
    set_run_font(run, size=10.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Tasker Consorcios")
    set_run_font(run, size=30, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("Documentación funcional, administrativa y técnica")
    set_run_font(run, size=14, color=DARK_BLUE)

    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(66)
    run = scope.add_run("Manual de usuario · Administración · Operación · API y datos")
    set_run_font(run, size=10.5, color=MUTED, italic=True)

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_after = Pt(4)
    run = version.add_run("Versión documentada 1.0 · Agosto de 2026")
    set_run_font(run, size=11, color=NAVY, bold=True)

    url = doc.add_paragraph()
    url.alignment = WD_ALIGN_PARAGRAPH.CENTER
    url.paragraph_format.space_after = Pt(3)
    run = url.add_run("tasker-consorcios.cuentagpt050.chatgpt.site")
    set_run_font(run, size=9.5, color=BLUE)

    repo = doc.add_paragraph()
    repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = repo.add_run("github.com/LeanBoveda/tasker-consorcios")
    set_run_font(run, size=9.5, color=BLUE)

    doc.add_page_break()


def add_front_matter(doc, bullet_abstract_id):
    heading = doc.add_paragraph(style="Heading 1")
    add_inline_runs(heading, "Cómo usar esta documentación", size=16, color=BLUE)
    for run in heading.runs:
        run.bold = True

    paragraph = doc.add_paragraph()
    add_inline_runs(paragraph, "Este documento consolida los manuales operativos y la referencia técnica de la versión productiva de Tasker Consorcios. Está pensado para usuarios, administradores y responsables de mantenimiento.")

    note = doc.add_paragraph()
    note.paragraph_format.left_indent = Inches(0.18)
    note.paragraph_format.right_indent = Inches(0.18)
    note.paragraph_format.space_before = Pt(6)
    note.paragraph_format.space_after = Pt(12)
    set_paragraph_shading(note, LIGHT_BLUE)
    add_inline_runs(note, "Importante: GitHub conserva el código y esta documentación. Las tareas, usuarios, consorcios y comentarios de producción se guardan por separado en Cloudflare D1.", color=NAVY)

    heading = doc.add_paragraph(style="Heading 2")
    add_inline_runs(heading, "Contenido", size=13, color=BLUE)
    for run in heading.runs:
        run.bold = True

    contents_num_id = create_num_instance(doc, bullet_abstract_id)
    for item in [
        "Manual de usuario",
        "Manual del administrador",
        "Documentación técnica",
        "Operación, mantenimiento y datos",
        "Referencia de API y modelo de datos",
    ]:
        paragraph = doc.add_paragraph()
        apply_numbering(paragraph, contents_num_id)
        add_inline_runs(paragraph, item)


def build():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Documentación completa de Tasker Consorcios"
    doc.core_properties.subject = "Manual funcional, administrativo y técnico"
    doc.core_properties.author = "Tasker Consorcios"
    doc.core_properties.keywords = "Tasker, consorcios, tareas, administración, manual"

    bullet_abstract_id = create_abstract_numbering(doc, bullet=True)
    decimal_abstract_id = create_abstract_numbering(doc, bullet=False)

    add_cover(doc)
    add_front_matter(doc, bullet_abstract_id)

    for source in SOURCE_FILES:
        doc.add_page_break()
        add_markdown(doc, source.read_text(encoding="utf-8"), bullet_abstract_id, decimal_abstract_id)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
